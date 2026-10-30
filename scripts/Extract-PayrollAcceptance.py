"""Read-only extraction of the user's two payroll workbooks for regression tests."""
import hashlib
import json
from pathlib import Path
import openpyxl

root = Path(__file__).resolve().parents[1]
samples = []
sources = []
for filename in ['5-Taşıma Bordro GENEL (1).xlsx', '5-Taşıma Bordro.xlsx']:
    path = Path.home() / 'Downloads' / filename
    values = openpyxl.load_workbook(path, data_only=True)
    formulas = openpyxl.load_workbook(path, data_only=False)
    sources.append({'File': filename, 'Sha256': hashlib.sha256(path.read_bytes()).hexdigest()})
    for index, sheet in enumerate(values):
        for row in range(2, sheet.max_row + 1):
            def cell(col): return sheet[f'{col}{row}'].value
            def num(col): return cell(col) or 0
            if not isinstance(cell('E' if index == 0 else 'D'), (int, float)) or not cell('B'):
                continue
            data = dict(WorkDays=num('E'), DailyPrice=num('F'), DiscountAmount=num('H'), PriceDifference=num('J'),
                        ContractAmount=0, IsVatExempt=num('K') == 0, TaxDebt=num('R'), SocialSecurityDebt=num('S'),
                        MinimumWageSupport=0, InspectionPenalty=num('Q'), ProgramType=1) if index == 0 else dict(
                        WorkDays=num('D'), DailyPrice=num('E'), DiscountAmount=0, PriceDifference=num('J'),
                        ContractAmount=0, IsVatExempt=num('M') == 0, TaxDebt=0, SocialSecurityDebt=0,
                        MinimumWageSupport=num('T'), InspectionPenalty=num('U'), GuideWorkDays=num('F'),
                        GuideDailyPrice=num('G'), GuideDiscountAmount=0, GuidePriceDifference=num('K'), ProgramType=0)
            mapping = {'G':'ServiceAmount', 'I':'TotalAmount', 'K':'Vat', 'L':'Withholding',
                       'M':'InvoiceAmountAfterWithholding', 'N':'GrossAmount', 'O':'StampTax',
                       'P':'AdditionalGuarantee', 'T':'TotalDeductions', 'U':'NetPayable'} if index == 0 else {
                       'H':'ServiceAmount', 'I':'GuideServiceAmount', 'L':'TotalAmount', 'M':'TransportVat',
                       'N':'TransportWithholding', 'O':'GuideVat', 'P':'GuideWithholding', 'Q':'GrossAmount',
                       'R':'AdditionalGuarantee', 'S':'StampTax', 'V':'NetPayable'}
            expected = []
            for col, prop in mapping.items():
                address = f'{col}{row}'
                expected.append(dict(Cell=address, Property=prop, Value=num(col),
                                     Formula=formulas[sheet.title][address].value,
                                     ReferenceMissing=index == 0 and col == 'M' and row == 21,
                                     SourceInvoiceVariant=index == 0 and col == 'M' and row == 23))
            samples.append(dict(File=filename, Sheet=sheet.title, Row=row, Input=data, Expected=expected))
out = root / 'tests/OzelEgitimHakedis.Tests/Fixtures/payroll-2026-10-26.json'
out.parent.mkdir(parents=True, exist_ok=True)
out.write_text(json.dumps(dict(Sources=sources, Samples=samples), ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps(dict(Rows=len(samples), Fields=sum(len(s['Expected']) for s in samples), Sources=sources), ensure_ascii=True))
